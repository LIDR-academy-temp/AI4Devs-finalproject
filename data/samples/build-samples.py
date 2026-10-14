#!/usr/bin/env python3
"""Converts plan.json into the sample .xlsx files (standardized template format)
and the expected.json files. Requires openpyxl.

    python data/samples/build-samples.py

IDs are written as TEXT (cell format @) so that Excel does not convert them to
scientific notation, just as the team does when pasting into the template.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).resolve().parent
plan = json.loads((HERE / "plan.json").read_text(encoding="utf-8"))

NON_TRANSACTIONAL_FEES = [
    "Annual Fee", "Bram Fine", "Chargeback Assessment Fee", "Chargeback Order Deduction", "Chargeback over 1%",
    "Chargeback over 2%", "Chargeback over 3%", "Chargeback Reversed", "Cross Border Fee", "Customer Service Fee",
    "Monthly Fee", "Rebate 2%", "Rebate 3%", "Release of Frozen Funds", "Remittance Charges", "Service Fee",
    "Settlement Fee", "Setup Fee", "Special Adjustment", "Special Deduction", "Wire Fee",
]
CHB_ALERT_FEES = ["Chargeback Alert Fee", "Chargeback Alert Fee Ethoca", "Chargeback Alert Fee RDR"]
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")


def write_sheet(ws, columns, rows):
    ws.append(columns)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
    for r in rows:
        ws.append(r)
    # column A (ID / Merchant Order No# / Type of Fee) as text
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        for c in row:
            if c.value is not None:
                c.number_format = "@"
                c.value = str(c.value)
    # leftover =SUM cell below the last row when the template leaves it (empty ID)
    for i, r in enumerate(rows, start=2):
        if r and r[0] == "":
            ws.cell(row=i, column=2).value = f"=SUM(B2:B{i-1})"
            ws.cell(row=i, column=3).value = None
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 10


def build_file(folder: Path, f: dict, template: bool):
    wb = Workbook()
    wb.remove(wb.active)
    for name, sheet in f["sheets"].items():
        ws = wb.create_sheet(title=name)
        write_sheet(ws, sheet["columns"], sheet["rows"])
    if template:
        ws = wb.create_sheet(title="Settings")
        ws["A1"] = "Chargeback Alert Fee Types"
        for i, v in enumerate(CHB_ALERT_FEES, start=3):
            ws.cell(row=i, column=1).value = v
        ws["C1"] = "Non-Transactional Fees"
        for i, v in enumerate(NON_TRANSACTIONAL_FEES, start=3):
            ws.cell(row=i, column=3).value = v
        ws.sheet_state = "hidden"
    out = folder / f["fileName"]
    wb.save(out)
    return out


def main():
    for s in plan["samples"]:
        folder = HERE / s["name"]
        folder.mkdir(exist_ok=True)
        template = s["name"] != "legacy-provider"
        for f in s["files"]:
            out = build_file(folder, f, template)
            print("wrote", out.relative_to(HERE))
        expected = {
            "sample": s["name"],
            "description": s["description"],
            "account": s["account"],
            "files": [f["fileName"] for f in s["files"]],
            **s["expected"],
        }
        (folder / "expected.json").write_text(json.dumps(expected, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        print("wrote", (folder / "expected.json").relative_to(HERE))


if __name__ == "__main__":
    sys.exit(main())
